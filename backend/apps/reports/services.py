from django.utils import timezone
from django.db.models import Count, Sum, Avg, Q
from datetime import timedelta, datetime
from apps.books.models import Book, Category, Rating
from apps.borrowing.models import BorrowRecord
from apps.users.models import User
from apps.reports.models import Report
import json
import csv
import io
import logging

logger = logging.getLogger(__name__)

# Try to import optional packages
try:
    from openpyxl import Workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    logger.warning("openpyxl not installed. Excel export disabled.")

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False
    logger.warning("reportlab not installed. PDF export disabled.")

class ReportService:
    
    # ==================== Core Report Generation ====================
    
    @staticmethod
    def generate_report(report_type, start_date=None, end_date=None, parameters=None):
        """Generate report data based on type"""
        if not start_date:
            start_date = timezone.now().date() - timedelta(days=30)
        if not end_date:
            end_date = timezone.now().date()
        
        # Convert string dates to date objects if needed
        if isinstance(start_date, str):
            start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
        if isinstance(end_date, str):
            end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
        
        report_functions = {
            'borrowing': ReportService._generate_borrowing_report,
            'user': ReportService._generate_user_report,
            'book': ReportService._generate_book_report,
            'overdue': ReportService._generate_overdue_report,
            'fine': ReportService._generate_fine_report,
            'inventory': ReportService._generate_inventory_report,
            'popular_books': ReportService._generate_popular_books_report,
            'category': ReportService._generate_category_report,
            'damage': ReportService._generate_damage_report,
            'digital_resources': ReportService._generate_digital_resources_report,
            'daily': ReportService._generate_daily_report,
            'monthly': ReportService._generate_monthly_report,
            'yearly': ReportService._generate_yearly_report,
        }
        
        func = report_functions.get(report_type)
        if func:
            return func(start_date, end_date, parameters)
        return None
    
    # ==================== Report Generation Functions ====================
    
    @staticmethod
    def _generate_borrowing_report(start_date, end_date, parameters=None):
        """Generate borrowing report"""
        borrowings = BorrowRecord.objects.filter(
            borrow_date__date__gte=start_date,
            borrow_date__date__lte=end_date
        )
        
        data = {
            'total_borrowings': borrowings.count(),
            'active_borrowings': borrowings.filter(status__in=['borrowed', 'renewed']).count(),
            'returned_borrowings': borrowings.filter(status='returned').count(),
            'overdue_borrowings': borrowings.filter(status='overdue').count(),
            'borrowings_by_status': {},
            'borrowings_by_role': {},
            'borrowings_by_day': {},
            'average_borrow_duration': borrowings.aggregate(Avg('borrow_duration_days'))['borrow_duration_days__avg'] or 0,
            'total_fines': float(borrowings.aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'recent_borrowings': [
                {
                    'user': b.user.username,
                    'book': b.book.title,
                    'borrow_date': b.borrow_date.strftime('%Y-%m-%d'),
                    'due_date': b.due_date.strftime('%Y-%m-%d'),
                    'status': b.status,
                    'fine': float(b.fine_amount)
                }
                for b in borrowings.order_by('-borrow_date')[:50]
            ]
        }
        
        # Group by status
        for status, _ in BorrowRecord.STATUS_CHOICES:
            count = borrowings.filter(status=status).count()
            if count > 0:
                data['borrowings_by_status'][status] = count
        
        # Group by role
        for role, _ in BorrowRecord.ROLE_CHOICES:
            count = borrowings.filter(borrower_role=role).count()
            if count > 0:
                data['borrowings_by_role'][role] = count
        
        # Group by day
        for i in range((end_date - start_date).days + 1):
            day = start_date + timedelta(days=i)
            count = borrowings.filter(borrow_date__date=day).count()
            if count > 0:
                data['borrowings_by_day'][day.strftime('%Y-%m-%d')] = count
        
        return data
    
    @staticmethod
    def _generate_user_report(start_date, end_date, parameters=None):
        """Generate user report"""
        users = User.objects.all()
        
        data = {
            'total_users': users.count(),
            'active_users': users.filter(is_active=True).count(),
            'verified_users': users.filter(is_email_verified=True).count(),
            'users_by_role': {},
            'new_users': users.filter(created_at__date__gte=start_date).count(),
            'users_with_borrowings': users.filter(borrowings__status__in=['borrowed', 'returned']).distinct().count(),
            'users_by_department': {},
            'recent_users': [
                {
                    'username': u.username,
                    'email': u.email,
                    'role': u.role,
                    'department': u.department or 'N/A',
                    'created_at': u.created_at.strftime('%Y-%m-%d'),
                    'borrowings_count': u.borrowings.count()
                }
                for u in users.order_by('-created_at')[:50]
            ]
        }
        
        # Group by role
        for role, _ in User.ROLE_CHOICES:
            count = users.filter(role=role).count()
            if count > 0:
                data['users_by_role'][role] = count
        
        # Group by department
        departments = users.values('department').annotate(count=Count('id')).order_by('-count')
        for dept in departments:
            if dept['department']:
                data['users_by_department'][dept['department']] = dept['count']
        
        return data
    
    @staticmethod
    def _generate_book_report(start_date, end_date, parameters=None):
        """Generate book report"""
        books = Book.objects.all()
        
        data = {
            'total_books': books.count(),
            'available_books': books.filter(available_quantity__gt=0).count(),
            'borrowed_books': books.filter(status='borrowed').count(),
            'total_copies': books.aggregate(Sum('quantity'))['quantity__sum'] or 0,
            'available_copies': books.aggregate(Sum('available_quantity'))['available_quantity__sum'] or 0,
            'books_by_category': {},
            'books_by_status': {},
            'most_borrowed_books': [],
            'books_with_reviews': books.filter(reviews__isnull=False).distinct().count(),
            'average_rating': float(Rating.objects.aggregate(Avg('rating'))['rating__avg'] or 0),
            'top_rated_books': [],
            'recent_books': [
                {
                    'title': b.title,
                    'author': b.author,
                    'category': b.category.name if b.category else 'N/A',
                    'status': b.status,
                    'available_quantity': b.available_quantity,
                    'created_at': b.created_at.strftime('%Y-%m-%d')
                }
                for b in books.order_by('-created_at')[:50]
            ]
        }
        
        # Group by category
        categories = books.values('category__name').annotate(count=Count('id')).order_by('-count')
        for cat in categories:
            if cat['category__name']:
                data['books_by_category'][cat['category__name']] = cat['count']
        
        # Group by status
        for status, _ in Book.STATUS_CHOICES:
            count = books.filter(status=status).count()
            if count > 0:
                data['books_by_status'][status] = count
        
        # Most borrowed books
        most_borrowed = books.annotate(
            borrow_count=Count('borrowings')
        ).filter(borrow_count__gt=0).order_by('-borrow_count')[:10]
        
        for b in most_borrowed:
            data['most_borrowed_books'].append({
                'title': b.title,
                'author': b.author,
                'borrow_count': b.borrowings.count(),
                'average_rating': float(b.reviews.aggregate(Avg('rating'))['rating__avg'] or 0)
            })
        
        # Top rated books
        top_rated = books.annotate(
            avg_rating=Avg('reviews__rating')
        ).filter(avg_rating__isnull=False).order_by('-avg_rating')[:10]
        
        for b in top_rated:
            data['top_rated_books'].append({
                'title': b.title,
                'author': b.author,
                'average_rating': float(b.reviews.aggregate(Avg('rating'))['rating__avg'] or 0),
                'review_count': b.reviews.count()
            })
        
        return data
    
    @staticmethod
    def _generate_overdue_report(start_date, end_date, parameters=None):
        """Generate overdue report"""
        overdue = BorrowRecord.objects.filter(status='overdue')
        
        data = {
            'total_overdue': overdue.count(),
            'total_fines': float(overdue.aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'overdue_by_user': {},
            'overdue_by_book': {},
            'overdue_list': [
                {
                    'user': b.user.username,
                    'user_id': b.get_borrower_id(),
                    'book': b.book.title,
                    'borrow_date': b.borrow_date.strftime('%Y-%m-%d'),
                    'due_date': b.due_date.strftime('%Y-%m-%d'),
                    'days_overdue': (timezone.now().date() - b.due_date).days,
                    'fine_amount': float(b.fine_amount),
                    'borrower_name': b.get_borrower_name(),
                    'borrower_id': b.get_borrower_id()
                }
                for b in overdue.order_by('due_date')[:50]
            ]
        }
        
        # Group by user
        overdue_users = overdue.values('user__username').annotate(count=Count('id')).order_by('-count')
        for u in overdue_users:
            data['overdue_by_user'][u['user__username']] = u['count']
        
        # Group by book
        overdue_books = overdue.values('book__title').annotate(count=Count('id')).order_by('-count')
        for b in overdue_books:
            data['overdue_by_book'][b['book__title']] = b['count']
        
        return data
    
    @staticmethod
    def _generate_fine_report(start_date, end_date, parameters=None):
        """Generate fine report"""
        fines = BorrowRecord.objects.filter(fine_amount__gt=0)
        
        data = {
            'total_fines': fines.count(),
            'total_amount': float(fines.aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'paid_fines': fines.filter(fine_paid=True).count(),
            'unpaid_fines': fines.filter(fine_paid=False).count(),
            'paid_amount': float(fines.filter(fine_paid=True).aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'unpaid_amount': float(fines.filter(fine_paid=False).aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'fines_by_user': {},
            'fines_by_month': {},
            'fine_list': [
                {
                    'user': b.user.username,
                    'book': b.book.title,
                    'fine_amount': float(b.fine_amount),
                    'status': 'Paid' if b.fine_paid else 'Unpaid',
                    'due_date': b.due_date.strftime('%Y-%m-%d'),
                    'return_date': b.return_date.strftime('%Y-%m-%d') if b.return_date else 'N/A'
                }
                for b in fines.order_by('-fine_amount')[:50]
            ]
        }
        
        # Group by user
        fine_users = fines.values('user__username').annotate(total=Sum('fine_amount')).order_by('-total')
        for u in fine_users:
            data['fines_by_user'][u['user__username']] = float(u['total'])
        
        # Group by month
        for i in range(12):
            month = timezone.now().date() - timedelta(days=30*i)
            month_name = month.strftime('%B %Y')
            month_fines = fines.filter(
                created_at__year=month.year,
                created_at__month=month.month
            ).aggregate(total=Sum('fine_amount'))['total'] or 0
            if month_fines > 0:
                data['fines_by_month'][month_name] = float(month_fines)
        
        return data
    
    @staticmethod
    def _generate_inventory_report(start_date, end_date, parameters=None):
        """Generate inventory report"""
        books = Book.objects.all()
        
        data = {
            'total_books': books.count(),
            'total_copies': books.aggregate(Sum('quantity'))['quantity__sum'] or 0,
            'available_copies': books.aggregate(Sum('available_quantity'))['available_quantity__sum'] or 0,
            'borrowed_copies': (books.aggregate(Sum('quantity'))['quantity__sum'] or 0) - (books.aggregate(Sum('available_quantity'))['available_quantity__sum'] or 0),
            'damaged_books': books.filter(status='damaged').count(),
            'lost_books': books.filter(status='lost').count(),
            'books_by_location': {},
            'books_by_category': {},
            'inventory_value': 0,
            'category_details': []
        }
        
        # Group by location
        locations = books.values('location').annotate(count=Count('id')).order_by('-count')
        for loc in locations:
            if loc['location']:
                data['books_by_location'][loc['location']] = loc['count']
        
        # Group by category with details
        categories = Category.objects.annotate(book_count=Count('books'))
        for cat in categories:
            cat_books = cat.books.all()
            data['category_details'].append({
                'name': cat.name,
                'total_books': cat_books.count(),
                'total_copies': cat_books.aggregate(Sum('quantity'))['quantity__sum'] or 0,
                'available_copies': cat_books.aggregate(Sum('available_quantity'))['available_quantity__sum'] or 0,
                'borrowed_copies': cat_books.filter(status='borrowed').count()
            })
        
        return data
    
    @staticmethod
    def _generate_popular_books_report(start_date, end_date, parameters=None):
        """Generate popular books report"""
        limit = parameters.get('limit', 10) if parameters else 10
        
        books = Book.objects.annotate(
            borrow_count=Count('borrowings', filter=Q(borrowings__status='returned')),
            avg_rating=Avg('reviews__rating'),
            review_count=Count('reviews')
        ).filter(borrow_count__gt=0).order_by('-borrow_count')[:limit]
        
        data = {
            'report_period': f"{start_date} to {end_date}",
            'popular_books': [
                {
                    'rank': i + 1,
                    'title': b.title,
                    'author': b.author,
                    'category': b.category.name if b.category else 'N/A',
                    'borrow_count': b.borrowings.filter(status='returned').count(),
                    'average_rating': float(b.reviews.aggregate(Avg('rating'))['rating__avg'] or 0),
                    'review_count': b.reviews.count(),
                    'available_copies': b.available_quantity
                }
                for i, b in enumerate(books)
            ]
        }
        
        return data
    
    @staticmethod
    def _generate_category_report(start_date, end_date, parameters=None):
        """Generate category report"""
        categories = Category.objects.annotate(
            book_count=Count('books'),
            borrow_count=Count('books__borrowings', filter=Q(books__borrowings__status='returned')),
            avg_rating=Avg('books__reviews__rating')
        ).order_by('-borrow_count')
        
        data = {
            'total_categories': categories.count(),
            'categories': [
                {
                    'name': cat.name,
                    'book_count': cat.book_count,
                    'borrow_count': cat.borrow_count,
                    'average_rating': float(cat.avg_rating or 0),
                    'popularity_score': cat.borrow_count * 0.7 + (cat.avg_rating or 0) * 0.3
                }
                for cat in categories
            ],
            'top_categories': []
        }
        
        # Top categories by borrow count
        top_cats = categories[:5]
        total_borrows = categories.aggregate(total=Sum('borrow_count'))['total'] or 1
        data['top_categories'] = [
            {
                'name': cat.name,
                'borrow_count': cat.borrow_count,
                'percentage': int(cat.borrow_count / total_borrows * 100)
            }
            for cat in top_cats
        ]
        
        return data
    
    @staticmethod
    def _generate_damage_report(start_date, end_date, parameters=None):
        """Generate damage report"""
        damaged = BorrowRecord.objects.filter(
            damage_status__in=['minor', 'moderate', 'severe', 'lost']
        )
        
        data = {
            'total_damaged': damaged.count(),
            'damage_by_status': {},
            'damage_costs': {},
            'damage_list': [
                {
                    'user': b.user.username,
                    'book': b.book.title,
                    'damage_status': b.damage_status,
                    'description': b.damage_description or 'N/A',
                    'reported_at': b.damage_reported_at.strftime('%Y-%m-%d') if b.damage_reported_at else 'N/A',
                    'cost': b.calculate_damage_cost()
                }
                for b in damaged.order_by('-damage_reported_at')[:50]
            ]
        }
        
        damage_costs_map = {
            'minor': 5,
            'moderate': 15,
            'severe': 30,
            'lost': 50
        }
        
        for status, label in BorrowRecord.DAMAGE_STATUS_CHOICES:
            count = damaged.filter(damage_status=status).count()
            if count > 0:
                data['damage_by_status'][status] = count
                data['damage_costs'][status] = count * damage_costs_map.get(status, 0)
        
        return data
    
    @staticmethod
    def _generate_digital_resources_report(start_date, end_date, parameters=None):
        """Generate digital resources report"""
        try:
            from apps.digital_resources.models import DigitalResource
            
            resources = DigitalResource.objects.all()
            
            data = {
                'total_resources': resources.count(),
                'resources_by_type': {},
                'resources_by_category': {},
                'recent_resources': [
                    {
                        'title': r.title,
                        'type': r.resource_type if hasattr(r, 'resource_type') else 'Unknown',
                        'category': r.category if hasattr(r, 'category') else 'Uncategorized',
                        'uploaded_at': r.uploaded_at.strftime('%Y-%m-%d') if hasattr(r, 'uploaded_at') else 'N/A',
                        'file_size': r.file_size if hasattr(r, 'file_size') else 0,
                        'downloads': r.download_count if hasattr(r, 'download_count') else 0
                    }
                    for r in resources.order_by('-uploaded_at')[:50]
                ]
            }
            
            # Group by type if field exists
            if hasattr(DigitalResource, 'resource_type'):
                resource_types = resources.values('resource_type').annotate(count=Count('id'))
                for rt in resource_types:
                    if rt['resource_type']:
                        data['resources_by_type'][rt['resource_type']] = rt['count']
            
            # Group by category if field exists
            if hasattr(DigitalResource, 'category'):
                categories = resources.values('category').annotate(count=Count('id'))
                for cat in categories:
                    if cat['category']:
                        data['resources_by_category'][cat['category']] = cat['count']
            
            return data
            
        except ImportError:
            return {
                'total_resources': 0,
                'message': 'Digital resources module not available'
            }
    
    @staticmethod
    def _generate_daily_report(start_date, end_date, parameters=None):
        """Generate daily report"""
        today = timezone.now().date()
        
        data = {
            'report_date': today.strftime('%Y-%m-%d'),
            'new_borrowings': BorrowRecord.objects.filter(borrow_date__date=today).count(),
            'new_returns': BorrowRecord.objects.filter(return_date=today).count(),
            'new_users': User.objects.filter(created_at__date=today).count(),
            'new_books': Book.objects.filter(created_at__date=today).count(),
            'overdue_count': BorrowRecord.objects.filter(status='overdue').count(),
            'active_borrowings': BorrowRecord.objects.filter(status__in=['borrowed', 'renewed']).count(),
            'visitors': 0,
            'daily_borrowings': [],
            'overdue_items': []
        }
        
        # Daily borrowings
        daily_borrowings = BorrowRecord.objects.filter(borrow_date__date=today)[:20]
        data['daily_borrowings'] = [
            {
                'user': b.user.username,
                'book': b.book.title,
                'borrow_date': b.borrow_date.strftime('%Y-%m-%d %H:%M'),
                'due_date': b.due_date.strftime('%Y-%m-%d')
            }
            for b in daily_borrowings
        ]
        
        # Overdue items
        overdue = BorrowRecord.objects.filter(status='overdue')[:20]
        data['overdue_items'] = [
            {
                'user': b.user.username,
                'book': b.book.title,
                'due_date': b.due_date.strftime('%Y-%m-%d'),
                'days': (today - b.due_date).days,
                'fine': float(b.fine_amount)
            }
            for b in overdue
        ]
        
        return data
    
    @staticmethod
    def _generate_monthly_report(start_date, end_date, parameters=None):
        """Generate monthly report"""
        today = timezone.now().date()
        month_start = today.replace(day=1)
        
        data = {
            'report_month': today.strftime('%B %Y'),
            'month_start': month_start.strftime('%Y-%m-%d'),
            'month_end': today.strftime('%Y-%m-%d'),
            'total_borrowings': BorrowRecord.objects.filter(borrow_date__date__gte=month_start).count(),
            'total_returns': BorrowRecord.objects.filter(return_date__gte=month_start).count(),
            'new_users': User.objects.filter(created_at__date__gte=month_start).count(),
            'new_books': Book.objects.filter(created_at__date__gte=month_start).count(),
            'total_fines_collected': float(BorrowRecord.objects.filter(
                fine_paid=True,
                return_date__gte=month_start
            ).aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'popular_books': [],
            'monthly_trend': []
        }
        
        # Popular books this month
        popular = Book.objects.annotate(
            borrow_count=Count('borrowings', filter=Q(borrowings__borrow_date__date__gte=month_start))
        ).filter(borrow_count__gt=0).order_by('-borrow_count')[:5]
        
        data['popular_books'] = [
            {
                'title': b.title,
                'author': b.author,
                'borrow_count': b.borrow_count
            }
            for b in popular
        ]
        
        # Monthly trend (last 12 months)
        for i in range(12):
            month = today.replace(day=1) - timedelta(days=30*i)
            count = BorrowRecord.objects.filter(
                borrow_date__year=month.year,
                borrow_date__month=month.month
            ).count()
            data['monthly_trend'].append({
                'month': month.strftime('%b %Y'),
                'borrowings': count
            })
        
        return data
    
    @staticmethod
    def _generate_yearly_report(start_date, end_date, parameters=None):
        """Generate yearly report"""
        year = timezone.now().year
        
        data = {
            'report_year': year,
            'total_borrowings': BorrowRecord.objects.filter(borrow_date__year=year).count(),
            'total_returns': BorrowRecord.objects.filter(return_date__year=year).count(),
            'new_users': User.objects.filter(created_at__year=year).count(),
            'new_books': Book.objects.filter(created_at__year=year).count(),
            'total_fines': float(BorrowRecord.objects.filter(
                created_at__year=year
            ).aggregate(Sum('fine_amount'))['fine_amount__sum'] or 0),
            'yearly_stats': [],
            'top_books_of_year': [],
            'monthly_breakdown': {}
        }
        
        # Monthly breakdown
        for month in range(1, 13):
            count = BorrowRecord.objects.filter(
                borrow_date__year=year,
                borrow_date__month=month
            ).count()
            data['monthly_breakdown'][month] = count
        
        # Top books of the year
        top_books = Book.objects.annotate(
            borrow_count=Count('borrowings', filter=Q(borrowings__borrow_date__year=year))
        ).filter(borrow_count__gt=0).order_by('-borrow_count')[:10]
        
        data['top_books_of_year'] = [
            {
                'title': b.title,
                'author': b.author,
                'borrow_count': b.borrow_count
            }
            for b in top_books
        ]
        
        return data
    
    # ==================== Report Export Functions ====================
    
    @staticmethod
    def export_to_pdf(report_data, title, filename=None):
        """Export report data to PDF"""
        if not HAS_REPORTLAB:
            return None
        
        try:
            if not filename:
                filename = f"report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            
            from django.http import HttpResponse
            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            
            doc = SimpleDocTemplate(response, pagesize=landscape(letter))
            styles = getSampleStyleSheet()
            story = []
            
            # Title
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                spaceAfter=30,
                alignment=1
            )
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 20))
            
            # Date
            date_style = ParagraphStyle(
                'DateStyle',
                parent=styles['Normal'],
                fontSize=10,
                textColor=colors.grey
            )
            story.append(Paragraph(f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}", date_style))
            story.append(Spacer(1, 20))
            
            # Format data for table
            if isinstance(report_data, dict):
                # Convert dict to table rows
                table_data = [['Metric', 'Value']]
                for key, value in report_data.items():
                    if not isinstance(value, (list, dict)):
                        table_data.append([str(key).replace('_', ' ').title(), str(value)])
                    elif isinstance(value, list) and len(value) > 0 and isinstance(value[0], dict):
                        # Handle list of dicts
                        headers = list(value[0].keys())
                        table_data = [['#'] + [h.replace('_', ' ').title() for h in headers]]
                        for i, item in enumerate(value):
                            row = [str(i+1)] + [str(item.get(h, '')) for h in headers]
                            table_data.append(row)
                        break
                
                # Create table
                if len(table_data) > 1:
                    col_widths = [1.5*inch] + [2*inch] * (len(table_data[0]) - 1)
                    table = Table(table_data, colWidths=col_widths[:len(table_data[0])])
                    table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 12),
                        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black)
                    ]))
                    story.append(table)
            
            doc.build(story)
            return response
            
        except Exception as e:
            logger.error(f"Error exporting to PDF: {str(e)}")
            return None
    
    @staticmethod
    def export_to_excel(report_data, title, filename=None):
        """Export report data to Excel"""
        if not HAS_OPENPYXL:
            return None
        
        try:
            if not filename:
                filename = f"report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            from django.http import HttpResponse
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            
            wb = Workbook()
            ws = wb.active
            ws.title = "Report"
            
            # Title
            ws['A1'] = title
            ws['A1'].font = Font(size=16, bold=True)
            ws.merge_cells('A1:D1')
            
            # Headers
            headers = ['Metric', 'Value', 'Details', 'Date']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            
            # Data
            row = 3
            if isinstance(report_data, dict):
                for key, value in report_data.items():
                    if not isinstance(value, (list, dict, tuple)):
                        ws.cell(row=row, column=1, value=str(key).replace('_', ' ').title())
                        ws.cell(row=row, column=2, value=str(value))
                        row += 1
            
            # Auto-adjust column widths
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(response)
            return response
            
        except Exception as e:
            logger.error(f"Error exporting to Excel: {str(e)}")
            return None
    
    @staticmethod
    def export_to_csv(report_data, title, filename=None):
        """Export report data to CSV"""
        try:
            import csv
            
            if not filename:
                filename = f"report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
            
            from django.http import HttpResponse
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            
            writer = csv.writer(response)
            
            # Write title
            writer.writerow([title])
            writer.writerow([])
            
            # Write data
            if isinstance(report_data, dict):
                writer.writerow(['Metric', 'Value'])
                for key, value in report_data.items():
                    if not isinstance(value, (list, dict, tuple)):
                        writer.writerow([str(key).replace('_', ' ').title(), str(value)])
            elif isinstance(report_data, list) and len(report_data) > 0 and isinstance(report_data[0], dict):
                # Write headers
                headers = list(report_data[0].keys())
                writer.writerow([h.replace('_', ' ').title() for h in headers])
                
                # Write rows
                for item in report_data:
                    writer.writerow([str(item.get(h, '')) for h in headers])
            
            return response
            
        except Exception as e:
            logger.error(f"Error exporting to CSV: {str(e)}")
            return None